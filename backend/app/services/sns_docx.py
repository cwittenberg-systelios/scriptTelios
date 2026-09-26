"""
DOCX-Export der SNS-Verlaufsauswertung (v19.41, Spec Abschnitt 7).

A4, Raender 2 cm, Kopfzeile "Vertraulich – Patientendaten · <Kürzel>", Seitenzahlen,
Ueberschriften in Klinik-Rot #971321, Tabellen nativ, Abbildungen nummeriert.
Eingabe: das Job-Ergebnis (text, fakten, grafiken, phasen) + Kuerzel/Anrede;
der Text darf im Frontend editiert worden sein (F2=B).
"""
from __future__ import annotations

import base64
import datetime as dt
import io
import re

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from app.services.sns_llm import ABSCHNITTE

BRAND = RGBColor(0x97, 0x13, 0x21)
GREY = RGBColor(0x52, 0x51, 0x4E)
FIG_WIDTH = Cm(17)


def _d(iso: str | None) -> str:
    if not iso:
        return "–"
    x = dt.date.fromisoformat(iso)
    return f"{x.day:02d}.{x.month:02d}.{x.year}"


def _f(v, nd=1) -> str:
    if v is None:
        return "–"
    if isinstance(v, float):
        return f"{v:.{nd}f}".replace(".", ",")
    return str(v)


def _field(run, instr: str):
    for tag, text in (("w:fldChar", "begin"), ("w:instrText", instr), ("w:fldChar", "end")):
        el = OxmlElement(tag)
        if tag == "w:fldChar":
            el.set(qn("w:fldCharType"), text)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def _shade(cell, hex_fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _table(doc, header: list[str], rows: list[list[str]], widths: list[float] | None = None,
           caption: str | None = None):
    if caption:
        p = doc.add_paragraph()
        r = p.add_run(caption)
        r.bold = True
        r.font.size = Pt(9)
        r.font.color.rgb = GREY
        p.paragraph_format.space_after = Pt(2)
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(8.5)
        _shade(c, "F2E4E6")
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(v))
            run.font.size = Pt(8.5)
    if widths:
        t.autofit = False
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def _figure(doc, g: dict, nr: int):
    png = base64.b64decode(g["png_b64"])
    doc.add_picture(io.BytesIO(png), width=FIG_WIDTH)
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph()
    r = p.add_run(f"Abbildung {nr}: {g['titel']}")
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = GREY
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    if g.get("legende"):
        p.paragraph_format.space_after = Pt(2)
        lp = doc.add_paragraph()
        lr = lp.add_run("Schlüsselereignisse: " + "; ".join(g["legende"]))
        lr.font.size = Pt(8)
        lr.font.color.rgb = GREY
        lp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        lp.paragraph_format.space_after = Pt(8)


def split_sections(text: str) -> dict[int, str]:
    """Zerlegt den Berichtstext in die neun Abschnitte {nr: body}."""
    out: dict[int, str] = {}
    pattern = re.compile(r"(?:^|\n)\s*#*\s*(\d)\.\s+[^\n]*\n?", re.M)
    hits = list(pattern.finditer(text))
    if not hits:
        return {1: text.strip()}
    for i, m in enumerate(hits):
        nr = int(m.group(1))
        if not 1 <= nr <= 9:
            continue
        end = hits[i + 1].start() if i + 1 < len(hits) else len(text)
        body = text[m.end():end].strip()
        out[nr] = (out.get(nr, "") + "\n\n" + body).strip() if nr in out else body
    return out


def _heading(doc, text: str, level: int = 1):
    h = doc.add_heading(text, level=level)
    for r in h.runs:
        r.font.color.rgb = BRAND
        r.font.size = Pt(13 if level == 1 else 11)
    return h


def _body(doc, text: str):
    for para in re.split(r"\n\s*\n", text.strip()):
        para = para.strip()
        if not para:
            continue
        p = doc.add_paragraph()
        para = re.sub(r"\s*\n\s*", " ", para)
        # **fett** (Einstiege in Abschnitt 9) als eigene Runs
        for i, teil in enumerate(re.split(r"\*\*(.+?)\*\*", para)):
            if teil:
                p.add_run(teil).bold = bool(i % 2)
        p.paragraph_format.space_after = Pt(6)


def datenbasis(fakten: dict) -> str:
    """Zeile unter dem Titel: Datenbasis und Luecken (v19.41.3)."""
    z = fakten.get("zeitraum") or {}
    ism = fakten.get("ism") or {}
    teile = [f"HSF-Basisbogen ({z.get('messtage_hsf', '–')} Messtage)"]
    if z.get("messtage_ind"):
        ab = f" ab {_d(z['ind_start'])[:6]}" if z.get("ind_start") and z["ind_start"] != z.get("start") else ""
        name = f" „{ism['fragebogen_name']}“" if ism.get("fragebogen_name") else ""
        teile.append(f"individueller Fragebogen{name} ({len(ism.get('items') or [])} Items, "
                     f"{z['messtage_ind']} Messtage{ab})")
    teile.append(f"{z.get('tagebucheintraege', '–')} Tagebucheinträge")
    s = "Datenbasis: " + ", ".join(teile) + "."
    lk = fakten.get("luecken") or {}
    x = sorted(set(lk.get("hsf_x", [])))
    fehl = sorted(set(lk.get("hsf_fehlend", [])) - set(x))
    if x:
        s += f" Von SNS übertragen und als fehlend behandelt: {', '.join(_d(d) for d in x)}."
    if fehl:
        s += f" Ohne Messung: {', '.join(_d(d) for d in fehl)}."
    return s


def build_docx(result: dict, kuerzel: str, anrede: str = "Klientin") -> bytes:
    fakten = result.get("fakten") or {}
    grafiken = result.get("grafiken") or {}
    phasen = result.get("phasen") or fakten.get("phasen") or []
    text = result.get("text") or ""
    secs = split_sections(text)

    doc = Document()
    sec = doc.sections[0]
    sec.orientation = WD_ORIENT.PORTRAIT
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2))
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    hp = sec.header.paragraphs[0]
    hp.text = f"Vertraulich – Patientendaten · {kuerzel}"
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for r in hp.runs:
        r.font.size = Pt(8)
        r.font.color.rgb = GREY
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run("Seite ")
    r.font.size = Pt(8)
    r2 = fp.add_run()
    r2.font.size = Pt(8)
    _field(r2, "PAGE")
    r3 = fp.add_run(" von ")
    r3.font.size = Pt(8)
    r4 = fp.add_run()
    r4.font.size = Pt(8)
    _field(r4, "NUMPAGES")

    t = doc.add_heading(f"ISM-Auswertung – {kuerzel}", level=0)
    for r in t.runs:
        r.font.color.rgb = BRAND
        r.font.size = Pt(18)
    z = fakten.get("zeitraum") or {}
    sub = doc.add_paragraph()
    sr = sub.add_run(
        f"SNS-Prozessmonitoring {_d(z.get('start'))} – {_d(z.get('ende'))} · "
        f"{z.get('messtage_hsf', '–')} HSF-Messtage · erstellt {dt.date.today():%d.%m.%Y} · "
        f"{anrede} ({kuerzel})"
    )
    sr.font.size = Pt(9)
    sr.font.color.rgb = GREY
    sub.paragraph_format.space_after = Pt(2)
    db = doc.add_paragraph()
    dr = db.add_run(datenbasis(fakten))
    dr.font.size = Pt(8.5)
    dr.font.color.rgb = GREY
    db.paragraph_format.space_after = Pt(10)

    fig_nr = {k: v["nr"] for k, v in grafiken.items()}

    for nr, titel in enumerate(ABSCHNITTE, start=1):
        _heading(doc, titel)
        body = secs.get(nr, "")
        if nr == 2:
            _body(doc, body)
            hsf = fakten.get("hsf") or {}
            _table(doc, ["Faktor", "Dimension", "Items", "Polung", "Anfang", "Ende", "τ"],
                   [[f["name"], f["kurz"], ", ".join(f["items"]), f["polung"], _f(f["anfang"]),
                     _f(f["ende"]), _f(f["tau"], 2)] for f in hsf.get("faktoren", [])],
                   [1.2, 4.0, 6.2, 1.2, 1.4, 1.4, 1.2], "Tabelle 1: HSF-Basisbogen – Dimensionen und Faktoren")
            ism = fakten.get("ism")
            if ism:
                q = {"xml": "Fragebogen-XML", None: "ohne Zuordnung"}
                _table(doc, ["Faktor", "Item", "Polung"],
                       [[it.get("faktor") or "–", it["titel"], it["polung"]] for it in ism.get("items", [])],
                       [1.6, 11.4, 4.0],
                       f"Tabelle 2: Individueller Fragebogen – Items und Faktorzuordnung ({q.get(ism.get('quelle'))})")
        elif nr == 3:
            _body(doc, body)
            if "hsf_faktoren" in grafiken:
                _figure(doc, grafiken["hsf_faktoren"], fig_nr["hsf_faktoren"])
        elif nr == 4:
            _body(doc, body)
            ism = fakten.get("ism")
            if ism and ism.get("faktoren"):
                labels = [ph["label"] for ph in phasen]
                _table(doc, ["Faktor", *labels, "Sprung", "Krisen-Min.", "Endniveau"],
                       [[f["name"], *[_f(f["phasenmittel"].get(lb)) for lb in labels],
                         _f(f["sprung_uebergang"]), _f(f["krisen_minimum"]), _f(f["endniveau"])]
                        if f["besetzt"] else [f["name"], *["–"] * len(labels), "unbesetzt", "–", "–"]
                        for f in ism["faktoren"]],
                       None, "Tabelle 3: ISM-Faktoren – Phasenmittel und Kennwerte (0–100)")
            if "ism_faktoren" in grafiken:
                _figure(doc, grafiken["ism_faktoren"], fig_nr["ism_faktoren"])
        elif nr == 5:
            _body(doc, body)
            for k in ("dk_resonanz", "krd"):
                if k in grafiken:
                    _figure(doc, grafiken[k], fig_nr[k])
        elif nr == 6:
            _body(doc, body)
            if "recurrence" in grafiken:
                _figure(doc, grafiken["recurrence"], fig_nr["recurrence"])
        elif nr == 7:
            _body(doc, body)
            gruppen = {g["key"]: g for g in (fakten.get("hsf") or {}).get("gruppen", [])}
            zfak = next((f for f in (fakten.get("ism") or {}).get("faktoren", []) if f["id"] == 0 and f["besetzt"]), None)
            extra = [("Symptome", gruppen.get("III")), ("Selbstwirks.", gruppen.get("VIII")), ("Zielerleben", zfak)]
            extra = [(n, g) for n, g in extra if g]
            _table(doc, ["Phase", "Zeitraum", "Name", *[n for n, _ in extra], "Komposit", "DK", "Res."],
                   [[ph["label"], f"{_d(ph['start'])[:6]}–{_d(ph['ende'])[:6]} ({ph['tage']} T.)", ph.get("name") or "–",
                     *[_f(g["phasenmittel"].get(ph["label"])) for _, g in extra],
                     _f(ph.get("komposit_mittel")), _f(ph.get("dk_mittel"), 3), ph.get("resonanz_max", "–")]
                    for ph in phasen], [1.1, 3.0, 4.2, *[1.6] * len(extra), 1.6, 1.3, 1.0],
                   "Tabelle 4: Phasen – Mittelwerte 0–100 (Symptome: niedriger = besser; Zielerleben = "
                   "ISM-Faktor I; DK = mittlere dynamische Komplexität; Res. = max. Anzahl kritischer Items)")
        elif nr == 8:
            _body(doc, body)
            rows = [[r["kurz"], r["polung"], _f(r["anfang"]), _f(r["ende"]), _f(r["delta"]), _f(r["tau"], 2)]
                    for r in (fakten.get("hsf") or {}).get("items", []) + fakten.get("ind_items", [])]
            _table(doc, ["Item", "Polung", "Anfang", "Ende", "Δ", "τ"], rows, [7.6, 1.4, 1.8, 1.8, 1.8, 1.6],
                   "Tabelle 5: Anfang und Ende je Item (Mittel der ersten/letzten drei Messtage)")
            if "hantel" in grafiken:
                _figure(doc, grafiken["hantel"], fig_nr["hantel"])
        else:
            _body(doc, body)

    hinweise = fakten.get("hinweise") or []
    if hinweise:
        _heading(doc, "Technische Hinweise", level=2)
        for h in hinweise:
            p = doc.add_paragraph(h)
            p.runs[0].font.size = Pt(8.5)
            p.runs[0].font.color.rgb = GREY
    par = fakten.get("parameter") or {}
    p = doc.add_paragraph()
    r = p.add_run(
        "Methodik: Dynamische Komplexität nach Schiepek & Strunk, identisch mit dem SNS (Fenster "
        f"{par.get('dk_fenster', 7)} Messtage, Wert am letzten Fenstertag; Kritikalität über dem "
        "95-%-Konfidenzintervall der Vorwerte); "
        f"Übergänge als Niveausprung im Ressourcen-Komposit (≥ {_f(par.get('major_shift', 20.0))} Punkte, "
        f"Niveauverschiebung ab {_f(par.get('min_shift', 10.0))}); Einbrüche ab {_f(par.get('einbruch_drop', 10.0))} "
        "Punkten unter dem Median der Vortage. Die Zeitreihen sind kurz "
        f"(n = {z.get('messtage_hsf', '–')} bzw. {z.get('messtage_ind', '–')}); Trends (Kendall τ) und "
        "Korrelationen sind deshalb deskriptiv zu lesen. Namen in Tagebuchzitaten sind durch Rollen ersetzt. "
        "Die Interpretation wurde KI-gestützt aus den berechneten "
        "Kennwerten formuliert und ist ärztlich-therapeutisch zu prüfen."
    )
    r.font.size = Pt(8)
    r.font.color.rgb = GREY

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
