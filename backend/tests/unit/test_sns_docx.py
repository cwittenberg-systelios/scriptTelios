"""Tests DOCX-Export + Endpoints /api/sns/docx, /api/sns/check (v19.41, S5)."""
from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.main import app
from app.services import sns_docx as sd
from app.services import sns_llm as sl
from app.services import sns_plots as sp
from app.services import sns_verlauf as sv

FIX = Path(__file__).resolve().parents[1] / "fixtures" / "sns_verlauf"


@pytest.fixture(scope="module")
def result():
    t = {n: (FIX / f).read_text(encoding="utf-8") for n, f in
         (("hsf", "hsf.csv"), ("ind", "individuell.csv"), ("xml", "individuell.xml"), ("doc", "faktor_I.doc"))}
    a = sv.analyse(t["hsf"], t["ind"], t["xml"], t["doc"])
    text = "\n\n".join(f"{h}\nAbsatz eins zu {h.split('. ', 1)[1]}.\n\nAbsatz zwei." for h in sl.ABSCHNITTE)
    text = text.replace("7. Phasen\n", "7. Phasen\nP1 – Ankommen: erste Tage.\n")
    phasen = sl.phasen_namen_anwenden(text, a.fakten["phasen"])
    return {"text": text, "fakten": a.fakten, "grafiken": sp.render_all(a), "phasen": phasen,
            "stage_a": [], "flags": a.flags, "namen": [], "quelle": [], "faktenblock": ""}


def test_split_sections():
    secs = sd.split_sections("1. Zusammenfassung\nA\n\n2. Fragebögen und Faktorstruktur\nB\n## 3. Verlauf\nC")
    assert secs == {1: "A", 2: "B", 3: "C"}


def test_build_docx_struktur(result):
    data = sd.build_docx(result, "FX12", "Klientin")
    doc = Document(io.BytesIO(data))
    heads = [p.text for p in doc.paragraphs if p.style.name.startswith(("Heading", "Title"))]
    assert heads[0] == "Verlaufsauswertung – FX12"
    for h in sl.ABSCHNITTE:
        assert h in heads
    assert len(doc.tables) == 5
    assert len(doc.inline_shapes) == 6
    captions = [p.text for p in doc.paragraphs if p.text.startswith("Abbildung ")]
    assert captions[0].startswith("Abbildung 1:") and captions[-1].startswith("Abbildung 6:")
    assert "FX12" in doc.sections[0].header.paragraphs[0].text
    assert doc.tables[3].rows[1].cells[3].text == "Ankommen"      # Phasenname aus Abschnitt 7
    assert doc.sections[0].left_margin.cm == pytest.approx(2.0, abs=0.01)


def test_build_docx_ohne_ism_und_grafiken(result):
    r = {**result, "fakten": {**result["fakten"], "ism": None, "ind_items": []}, "grafiken": {}}
    doc = Document(io.BytesIO(sd.build_docx(r, "K", "Klient")))
    assert len(doc.tables) == 3 and len(doc.inline_shapes) == 0


def test_endpoint_docx_und_check(result):
    client = TestClient(app)
    r = client.post("/api/sns/docx", json={"kuerzel": "FX 12/3", "anrede": "Klientin", "text": result["text"],
                                           "result": {k: v for k, v in result.items() if k != "text"}})
    assert r.status_code == 200, r.text
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="Verlaufsauswertung_FX123.docx"' in r.headers["content-disposition"]
    assert Document(io.BytesIO(r.content)).tables
    c = client.post("/api/sns/check", json={"text": result["text"].replace("6. Rekurrenzmuster", "6. x"),
                                            "result": {k: v for k, v in result.items() if k != "text"}})
    assert c.status_code == 200
    assert "SNS_ABSCHNITT_FEHLT" in [i["code"] for i in c.json()["issues"]]
    assert c.json()["summary"]["checks_run"] == 16
    assert client.post("/api/sns/docx", json={"kuerzel": "", "text": "x", "result": {}}).status_code == 422
    assert json.dumps(c.json())
